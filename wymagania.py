import io
from flask import Flask, make_response, jsonify
from flask_cors import CORS
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

app = Flask(__name__)
CORS(app)



MOCK_DATA = {
    "title": "Miesięczny Raport Sprzedaży",
    "date": "Sierpień 2026",
    "items": [
        {"id": 1, "name": "Usługa IT", "category": "Konsulting", "value": 5000.00},
        {"id": 2, "name": "Licencja SaaS", "category": "Oprogramowanie", "value": 1200.00},
        {"id": 3, "name": "Szkolenie zespołu", "category": "Edukacja", "value": 3500.00}
    ],
    "total": 9700.00
}

@app.route('/api/export/pdf', methods=['GET'])
def export_to_pdf():
    try:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, 
            pagesize=letter, 
            rightMargin=40, 
            leftMargin=40, 
            topMargin=40, 
            bottomMargin=40
        )
        story = []
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('Title', parent=styles['Heading1'], alignment=1, spaceAfter=10)
        date_style = ParagraphStyle('Date', parent=styles['Normal'], alignment=1, spaceAfter=25)
        
        story.append(Paragraph(MOCK_DATA["title"], title_style))
        story.append(Paragraph(f"Okres: {MOCK_DATA['date']}", date_style))
        
        table_data = [["Nazwa", "Kategoria", "Wartość (PLN)"]]
        for item in MOCK_DATA["items"]:
            table_data.append([item["name"], item["category"], f"{item['value']:.2f}"])
        table_data.append(["", "Suma:", f"{MOCK_DATA['total']:.2f}"])
        
        table = Table(table_data, colWidths=[230, 150, 150])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2C3E50')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('ALIGN', (0, 0), (1, -1), 'LEFT'),
            ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
            ('FONTNAME', (1, -1), (-1, -1), 'Helvetica-Bold'),
            ('LINEBELOW', (0, -2), (-1, -2), 1, colors.black),
            ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
        ]))
        
        story.append(table)
        doc.build(story)
        
        buffer.seek(0)
        response = make_response(buffer.getvalue())
        buffer.close()
        
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = 'attachment; filename=raport.pdf'
        return response

    except Exception as e:
        return jsonify({"error": f"Błąd PDF: {str(e)}"}), 500


@app.route('/api/export/excel', methods=['GET'])
def export_to_excel():
    try:
        buffer = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "Raport Sprzedaży"
        
        headers = ["ID", "Nazwa", "Kategoria", "Wartość (PLN)"]
        ws.append(headers)
        
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
            
        for item in MOCK_DATA["items"]:
            ws.append([item["id"], item["name"], item["category"], item["value"]])
            
        ws.append([])
        ws.append(["", "", "Suma:", MOCK_DATA["total"]])
        
        last_row = ws.max_row
        ws.cell(row=last_row, column=3).font = Font(bold=True)
        ws.cell(row=last_row, column=4).font = Font(bold=True)
        
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        wb.save(buffer)
        buffer.seek(0)
        response = make_response(buffer.getvalue())
        buffer.close()
        
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = 'attachment; filename=raport.xlsx'
        return response

    except Exception as e:
        return jsonify({"error": f"Błąd Excel: {str(e)}"}), 500


if __name__ == '__main__':
    import openpyxl
    app.run(port=5001, debug=True)
